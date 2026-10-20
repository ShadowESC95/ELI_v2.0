"""Markdown with LaTeX maths, typeset without a TeX installation.

matplotlib draws each formula (its mathtext reads the common LaTeX maths), Python-Markdown lays out the
text and Qt writes the PDF; all three ship in every ELI build, while pandoc and LuaLaTeX, which the
older export path needs, are on few machines. The chat, the task view and saved documents use the same
rendering, so a formula reads the same everywhere.
"""
from __future__ import annotations

import functools
import hashlib
import html
import io
import re
import struct
import threading
from pathlib import Path
from typing import Dict, List, Tuple

from eli.utils.log import get_logger

log = get_logger(__name__)

_DISPLAY = re.compile(r"\$\$(.+?)\$\$|\\\[(.+?)\\\]", re.S)
# $x$ but not prices: no space inside either dollar, no digit or letter right after the closing one
_INLINE = re.compile(r"\\\((.+?)\\\)|(?<![\\$\w])\$(?=\S)([^$\n]+?)(?<=\S)\$(?![\w$])")
_ENV = re.compile(r"\\begin\{(\w+)\*?\}(.*?)\\end\{\1\*?\}", re.S)
_CODE = re.compile(r"```.*?```|`[^`\n]+`", re.S)
_DRAW_LOCK = threading.Lock()      # matplotlib's text engine keeps global state
# LaTeX commands mathtext lacks, written with ones it has
_SPELLINGS = ((r"\Box", "\u25a1"), (r"\boldsymbol", r"\mathbf"), (r"\bm", r"\mathbf"),
              (r"\lvert", "|"), (r"\rvert", "|"), (r"\lVert", r"\Vert"), (r"\rVert", r"\Vert"),
              (r"\coloneqq", ":="), (r"\tfrac", r"\frac"), (r"\dfrac", r"\frac"))


def has_markup(text: str) -> bool:
    """Formulas, or Markdown worth laying out (headings, emphasis, lists, tables)."""
    t = str(text or "")
    return bool(_DISPLAY.search(t) or _INLINE.search(t)
                or re.search(r"(?m)^\s{0,3}(?:#{1,6}\s|[-*+]\s|\d+[.)]\s|\|.*\|)|\*\*[^*\n]+\*\*", t))


@functools.lru_cache(maxsize=None)
def formula_png(tex: str, *, size_pt: float, dpi: float, color: str) -> bytes:
    """One formula as a PNG, or b"" when matplotlib cannot read it."""
    for latex, mathtext_name in _SPELLINGS:
        tex = re.sub(re.escape(latex) + r"(?![A-Za-z])", lambda _m: mathtext_name, tex)
    try:
        import matplotlib
        from matplotlib import mathtext
        from matplotlib.font_manager import FontProperties
        buf = io.BytesIO()
        with _DRAW_LOCK, matplotlib.rc_context({"mathtext.fontset": "cm", "savefig.transparent": True}):
            mathtext.math_to_image(f"${tex}$", buf, prop=FontProperties(size=size_pt), dpi=dpi,
                                   format="png", color=color)
        return buf.getvalue()
    except Exception as exc:
        log.debug("formula left as source (%s): %s", str(exc).strip().splitlines()[-1][:120], tex[:80])
        return b""


def _png_size(png: bytes) -> Tuple[int, int]:
    return struct.unpack(">II", png[16:24])


def _rows(tex: str) -> List[str]:
    """An aligned block (aligned, align, gather, equation, split) one row at a time: mathtext reads
    single expressions only."""
    m = _ENV.search(tex)
    if not m:
        return [tex]
    rows = [r.replace("&", " ").strip() for r in re.split(r"\\\\", m.group(2))]
    return [r for r in rows if r]


def to_html(md: str, *, color: str, size_pt: float, scale: float) -> Tuple[str, Dict[str, bytes]]:
    """HTML for a Qt text document, and the formula images it names (url -> PNG).

    `scale` is device pixels per layout pixel (a 96 dpi page), so formulas stay sharp on the screen
    or page they are drawn for and sit at the text's size."""
    images: Dict[str, bytes] = {}
    slots: List[str] = []
    codes: List[str] = []

    def stash_code(m):
        codes.append(m.group(0))
        return f"ELICODESLOT{len(codes) - 1}X"

    def formula(tex: str) -> str:
        png = formula_png(tex.strip(), size_pt=size_pt, dpi=96.0 * scale, color=color)
        if not png:
            return f"<code>{html.escape(tex.strip())}</code>"
        url = "eli-formula:" + hashlib.sha1(png).hexdigest()[:16]
        images[url] = png
        w, h = _png_size(png)
        return (f'<img src="{url}" width="{max(1, round(w / scale))}" '
                f'height="{max(1, round(h / scale))}" style="vertical-align: middle">')

    def display(m):
        tex = m.group(1) or m.group(2) or ""
        slots.append("<br>".join(formula(r) for r in _rows(tex)))
        return f"\n\nELIMATHSLOT{len(slots) - 1}D\n\n"

    def inline(m):
        slots.append(formula(m.group(1) or m.group(2) or ""))
        return f"ELIMATHSLOT{len(slots) - 1}X"

    text = _CODE.sub(stash_code, str(md or ""))
    text = _DISPLAY.sub(display, text)
    text = _INLINE.sub(inline, text)
    text = _nest_lists(text)
    # Markdown starts a list only after a blank line; replies often put one straight under a sentence
    text = re.sub(r"(?m)^(?!\s*(?:[-*+]|\d+[.)])\s)(\S.*)\n(?=\s{0,3}(?:[-*+]|\d+[.)])\s)", r"\1\n\n", text)
    text = re.sub(r"ELICODESLOT(\d+)X", lambda m: codes[int(m.group(1))], text)
    try:
        import markdown
        body = markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
    except Exception:
        log.debug("markdown unavailable; plain paragraphs", exc_info=True)
        body = "".join(f"<p>{html.escape(p).replace(chr(10), '<br>')}</p>" for p in text.split("\n\n"))
    body = body.replace("<table>", '<table border="1" cellspacing="0" cellpadding="4">')
    body = re.sub(r"<p>\s*ELIMATHSLOT(\d+)D\s*</p>",
                  lambda m: f'<p align="center">{slots[int(m.group(1))]}</p>', body)
    body = re.sub(r"ELIMATHSLOT(\d+)[DX]", lambda m: slots[int(m.group(1))], body)
    return body, images


_ITEM = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s")


def _nest_lists(text: str) -> str:
    """Items indented two or three spaces under another item are its children; Markdown wants four
    and otherwise runs them into the parent's line."""
    out: List[str] = []
    in_list = False
    for line in text.split("\n"):
        m = re.match(r"^ {1,3}((?:[-*+]|\d+[.)])\s.*)$", line)
        if m and in_list:
            line = "    " + m.group(1)
        if line.strip():
            in_list = bool(_ITEM.match(line)) or (in_list and line.startswith(" "))
        out.append(line)
    return "\n".join(out)


def add_images(document, images: Dict[str, bytes]) -> None:
    """Hand the formula images to a QTextDocument before its HTML is set."""
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QImage, QTextDocument
    for url, png in images.items():
        document.addResource(QTextDocument.ResourceType.ImageResource, QUrl(url), QImage.fromData(png))


def _page_size():
    from PySide6.QtCore import QLocale
    from PySide6.QtGui import QPageSize
    us = QLocale.system().measurementSystem() == QLocale.MeasurementSystem.ImperialUSSystem
    return QPageSize(QPageSize.PageSizeId.Letter if us else QPageSize.PageSizeId.A4)


def write_pdf(md: str, path, *, title: str = "") -> Path:
    """The document as a PDF, laid out by Qt. Needs the Qt application of the ELI window."""
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QGuiApplication, QPageLayout, QPdfWriter, QTextDocument
    if QGuiApplication.instance() is None:
        raise RuntimeError("no Qt application in this process")
    path = Path(path)
    writer = QPdfWriter(str(path))
    writer.setPageSize(_page_size())
    layout = writer.pageLayout()
    if layout.margins(QPageLayout.Unit.Millimeter).isNull():
        side = layout.fullRect(QPageLayout.Unit.Millimeter).width() / 10   # a tenth of the width each side
        writer.setPageMargins(QMarginsF(side, side, side, side), QPageLayout.Unit.Millimeter)
    writer.setTitle(title or path.stem)
    writer.setCreator("ELI")
    doc = QTextDocument()
    size = doc.defaultFont().pointSizeF()
    body, images = to_html(md, color="#000000", size_pt=size, scale=writer.resolution() / 96.0)
    add_images(doc, images)
    head = f"<h1>{html.escape(title)}</h1>" if title else ""
    doc.setHtml(head + body)
    doc.print_(writer)
    return path


def write_docx(md: str, path, *, title: str = "") -> Path:
    """The document as .docx: headings, lists and paragraphs, formulas as pictures."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    path = Path(path)
    doc = Document()
    size = doc.styles["Normal"].font.size or Pt(11)
    if title:
        doc.add_heading(title, level=0)

    def runs(par, text: str) -> None:
        pos = 0
        for m in _INLINE.finditer(text):
            _words(par, text[pos:m.start()])
            png = formula_png((m.group(1) or m.group(2) or "").strip(), size_pt=size.pt, dpi=300.0, color="#000000")
            if png:
                par.add_run().add_picture(io.BytesIO(png), height=Pt(_png_size(png)[1] * 72.0 / 300.0))
            else:
                par.add_run(m.group(0))
            pos = m.end()
        _words(par, text[pos:])

    text = _DISPLAY.sub(lambda m: "\n\nELIDISPLAY" + (m.group(1) or m.group(2) or "").replace("\n", " ") + "\n\n",
                        str(md or ""))
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        if block.startswith("ELIDISPLAY"):
            for row in _rows(block[len("ELIDISPLAY"):]):
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                png = formula_png(row, size_pt=size.pt, dpi=300.0, color="#000000")
                if png:
                    par.add_run().add_picture(io.BytesIO(png), height=Pt(_png_size(png)[1] * 72.0 / 300.0))
                else:
                    par.add_run(row)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", block)
        if heading and "\n" not in block:
            doc.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            continue
        for line in block.split("\n"):
            item = re.match(r"^\s*(?:([-*+])|(\d+)[.)])\s+(.*)$", line)
            if item:
                par = doc.add_paragraph(style="List Bullet" if item.group(1) else "List Number")
                runs(par, item.group(3))
            else:
                runs(doc.add_paragraph(), line.strip())
    doc.save(str(path))
    return path


def _words(par, text: str) -> None:
    """Text with **bold** and *italic*."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def write_document(md: str, path, *, title: str = "") -> Path:
    """PDF where Qt can lay it out, .docx otherwise (a server with no window); returns where it went."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        try:
            return write_pdf(md, path, title=title)
        except Exception:
            log.debug("PDF not written; writing .docx", exc_info=True)
            path = path.with_suffix(".docx")
    return write_docx(md, path, title=title)
